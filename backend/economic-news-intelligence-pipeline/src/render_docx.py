from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt


def render_daily_docx(
    config: dict[str, Any],
    target_date: date,
    domestic_items: list[dict[str, Any]],
    international_items: list[dict[str, Any]],
    overwrite: bool = False,
) -> Path:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(11)

    title = document.add_heading(f"{target_date.isoformat()} 每日新闻整理", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _add_section(document, f"中国国内新闻 Top {len(domestic_items)}", domestic_items)
    _add_section(document, f"国际新闻 Top {len(international_items)}", international_items)

    footer = document.add_paragraph()
    footer.add_run("Generated automatically by Kara's News Agent.")

    output_dir = Path(config["output_dir"])
    output_path = _resolve_output_path(output_dir, target_date, overwrite=overwrite)
    document.save(output_path)
    return output_path


def _add_section(document: Document, heading: str, items: list[dict[str, Any]]) -> None:
    document.add_heading(heading, level=1)
    if not items:
        document.add_paragraph("暂无符合条件的新闻。")
        return

    for index, item in enumerate(items, start=1):
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{index}. {item['title']}\n").bold = True
        paragraph.add_run(f"来源：{item['source']}\n")
        paragraph.add_run(f"发布日期：{item.get('published_date') or '未标明'}\n")
        paragraph.add_run(f"摘要：{_summary_text(item)}\n")
        link_url = item.get("final_url_used") or item["url"]
        paragraph.add_run("链接：")
        _append_hyperlink(paragraph, link_url, link_url)
        if item.get("tags"):
            paragraph.add_run(f"\n标签：{' / '.join(item['tags'])}")
        if item.get("reason"):
            paragraph.add_run(f"\n入选原因：{item['reason']}")


def _summary_text(item: dict[str, Any]) -> str:
    summary_zh = (item.get("summary_zh") or "").strip()
    if summary_zh and item.get("summary_source") == "llm":
        return summary_zh
    return ""


def _resolve_output_path(output_dir: Path, target_date: date, overwrite: bool) -> Path:
    base_path = output_dir / f"{target_date.isoformat()}_新闻整理.docx"
    if overwrite or not base_path.exists():
        return base_path

    version_path = output_dir / f"{target_date.isoformat()}_新闻整理_v2.docx"
    if not version_path.exists():
        return version_path

    timestamp = datetime.now().strftime("%H%M%S")
    return output_dir / f"{target_date.isoformat()}_新闻整理_{timestamp}.docx"


def _append_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)

    new_run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)

    new_run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    new_run.append(text_element)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
